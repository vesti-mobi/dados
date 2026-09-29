"""
fetch_churn_planilha.py — churn declarado pelo time, da planilha do Google.

Pedido do Walid em 29/09/2026: "puxe os churn da planilha ... aba 2026 com a
coluna observacao escrita solicitou cancelamento e inadimplente, os outros
ignore". O painel já tem duas leituras de churn que saem de fato (a marca parou
de pagar) — esta é a terceira e é diferente das outras: é o que o time
REGISTROU, com o motivo escrito por uma pessoa.

Saída: CS/churn_planilha.json, que o fetch_dados.js embute no dados.js. O
arquivo é intermediário e não vai para o git.

Acesso, em duas tentativas, nesta ordem:
  1. download público (planilha com "qualquer pessoa com o link pode ver");
  2. a service account do BigQuery (GOOGLE_APPLICATION_CREDENTIALS, que no
     Actions vem do secret GCP_SA_KEY) — para isso a planilha precisa estar
     compartilhada com o e-mail dela como Leitor.
Sem nenhum dos dois, o script avisa e sai sem escrever: a carga segue e só o
card do churn declarado fica vazio.

Rodar:  python3 fetch_churn_planilha.py
"""

from __future__ import annotations  # anotacoes `X | None` em Python 3.9

import csv
import io
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(__file__).parent
OUT = ROOT / "churn_planilha.json"
SA_JSON = Path(os.environ.get("GOOGLE_APPLICATION_CREDENTIALS") or (ROOT / "google_sa.json"))

SHEET_ID = os.environ.get("CHURN_SHEET_ID") or "19gGUPtmDT9xacgfkTWcLM_MN_iWhUGJQOjFaD9LXasQ"
ABA = os.environ.get("CHURN_SHEET_ABA") or "2026"

# Só estas observações contam como churn. O resto da planilha é ignorado —
# regra explícita do Walid, não invenção nossa.
OBSERVACOES = ["solicitou cancelamento", "inadimplente"]

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]


def chave(texto: str) -> str:
    """Compara sem acento, sem caixa e sem espaço sobrando."""
    t = unicodedata.normalize("NFD", str(texto or ""))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", t).strip().lower()


def baixar(url: str, token: str | None = None) -> bytes:
    cab = {"User-Agent": "Mozilla/5.0"}
    if token:
        cab["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, headers=cab)
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()


def token_sa() -> str:
    if not SA_JSON.exists():
        raise RuntimeError(
            f"credencial nao encontrada em {SA_JSON} — aponte "
            "GOOGLE_APPLICATION_CREDENTIALS para a chave da service account"
        )
    from google.oauth2 import service_account
    import google.auth.transport.requests

    creds = service_account.Credentials.from_service_account_file(str(SA_JSON), scopes=SCOPES)
    creds.refresh(google.auth.transport.requests.Request())
    email = getattr(creds, "service_account_email", "(desconhecido)")
    print(f"[churn] usando a service account {email}")
    return creds.token


def baixar_xlsx() -> bytes:
    """A planilha inteira, em xlsx — é o formato que traz TODAS as abas, e é
    delas que a de 2026 sai pelo nome, sem depender do gid."""
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    try:
        dados = baixar(url)
        if dados[:2] == b"PK":
            print("[churn] planilha baixada sem credencial (link público)")
            return dados
        print("[churn] o link público não devolveu xlsx; tentando com a service account")
    except urllib.error.HTTPError as e:
        print(f"[churn] link público respondeu HTTP {e.code}; tentando com a service account")

    dados = baixar(url, token_sa())
    if dados[:2] != b"PK":
        raise RuntimeError("a resposta não é um xlsx — a planilha não está compartilhada com a SA")
    return dados


def ler_aba(xlsx: bytes) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True, read_only=True)
    nome = next((n for n in wb.sheetnames if chave(n) == chave(ABA)), None)
    if not nome:
        raise RuntimeError(f'aba "{ABA}" não existe. Abas: {", ".join(wb.sheetnames)}')
    ws = wb[nome]
    linhas = list(ws.values)
    if not linhas:
        return []
    cabec = [chave(c) for c in linhas[0]]
    print(f'[churn] aba "{nome}": {len(linhas)-1} linhas · colunas: {", ".join(c for c in cabec if c)}')
    return [dict(zip(cabec, linha)) for linha in linhas[1:]]


# Fórmula quebrada na planilha chega como texto ("#REF!", "#N/D"); vale o mesmo
# que célula vazia.
ERRO_PLANILHA = re.compile(r"^#(ref|n/?d|name|value|div/0|null)", re.IGNORECASE)


def pegar(linha: dict, *nomes: str):
    for n in nomes:
        for k, v in linha.items():
            if k and (k == chave(n) or k.startswith(chave(n))):
                if v in (None, ""):
                    continue
                if isinstance(v, str) and ERRO_PLANILHA.match(v.strip()):
                    continue
                return v
    return None


def iso(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    t = str(v).strip()
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", t)
    if m:
        d, mes, ano = m.groups()
        ano = ("20" + ano) if len(ano) == 2 else ano
        return f"{ano}-{int(mes):02d}-{int(d):02d}"
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", t)
    return m.group(0) if m else None


def numero(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = re.sub(r"[^\d,.-]", "", str(v)).replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def main() -> None:
    try:
        linhas = ler_aba(baixar_xlsx())
    except Exception as e:  # noqa: BLE001 — qualquer falha aqui é "sem planilha hoje"
        print(f"[churn] não deu para ler a planilha: {e}", file=sys.stderr)
        print(
            "[churn] compartilhe a planilha com a service account (Leitor) ou "
            'deixe-a como "qualquer pessoa com o link pode ver". A carga segue '
            "sem o churn declarado.",
            file=sys.stderr,
        )
        sys.exit(1)

    aceitos, ignorados = [], {}
    for l in linhas:
        obs = pegar(l, "observacao", "observacoes", "obs")
        k = chave(obs)
        if not k:
            continue
        alvo = next((o for o in OBSERVACOES if k.startswith(o)), None)
        if not alvo:
            ignorados[k[:40]] = ignorados.get(k[:40], 0) + 1
            continue
        nome = pegar(l, "marca", "cliente", "empresa", "nome")
        if not nome:
            continue
        # A coluna de domínio da planilha vem de fórmula e chega como "#REF!" ou
        # "None" em boa parte das linhas — só vale quando é número mesmo.
        dom = re.sub(r"\D", "", str(pegar(l, "dominio", "domain") or ""))
        # `mensalidade` antes de `valor`: existe uma coluna "valor reajuste" que
        # o prefixo "valor" pegaria primeiro, e ela não é o que a marca pagava.
        valor = numero(pegar(l, "mensalidade", "total cobrado", "mrr", "valor"))
        aceitos.append({
            "cliente": str(nome).strip(),
            "dominio": dom or None,
            "cnpj": re.sub(r"\D", "", str(pegar(l, "cnpj", "cpf/cnpj") or "")) or None,
            "vendedora": (str(pegar(l, "vendedora", "vendedor", "cs", "anjo") or "")).strip() or None,
            "canal": (str(pegar(l, "canal") or "")).strip() or None,
            "subconta": (str(pegar(l, "subconta") or "")).strip() or None,
            "motivo": "Solicitou cancelamento" if alvo == "solicitou cancelamento" else "Inadimplente",
            "observacao": str(obs).strip(),
            "data": iso(pegar(l, "data da perda", "data", "data do churn", "mes", "cancelamento")),
            "plano": (str(pegar(l, "plano") or "")).strip() or None,
            "valor": valor,
        })

    print(f"[churn] aceitos: {len(aceitos)} (solicitou cancelamento / inadimplente)")
    if ignorados:
        top = sorted(ignorados.items(), key=lambda x: -x[1])[:6]
        print("[churn] observações ignoradas: " + ", ".join(f"{k} ({n})" for k, n in top))

    OUT.write_text(json.dumps({
        "aba": ABA,
        "observacoesAceitas": OBSERVACOES,
        "geradoEm": datetime.now().isoformat(timespec="seconds"),
        "linhas": aceitos,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"[write] {OUT.name}")


if __name__ == "__main__":
    main()
