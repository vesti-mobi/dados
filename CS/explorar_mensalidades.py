"""
explorar_mensalidades.py — só OLHA a planilha de mensalidades.

Pedido da Laura em 29/10/2026: quando uma marca aparece na Bonificação sem
valor porque não tem mensalidade no Iugu, puxar o valor desta planilha —
"sempre da aba vesti e starter e sempre olhando no mês mais atual que tem lá".

Antes de escrever o leitor de verdade é preciso saber como ela é por dentro:
que abas existem, que colunas cada uma tem e como os meses estão organizados.
É só isso que este script faz.

O LOG DO ACTIONS É PÚBLICO. Então aqui sai o ESQUELETO — nome de aba, nome de
coluna, contagem de linhas — e nunca o conteúdo: nem nome de marca, nem valor,
nem o e-mail da service account.

Roda no workflow "Conferir BigQuery" com modo=mensalidades.
"""

from __future__ import annotations

import io
import os
import sys
import urllib.error
import urllib.request
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(__file__).parent
# As credenciais possíveis, em ordem de preferência. A planilha de
# mensalidades foi compartilhada com a service account de leitura de
# planilhas, que não é a mesma do BigQuery — por isso a lista, e por isso o
# script diz QUAL funcionou (pelo apelido do secret, nunca pelo e-mail: o log
# deste repositório é público).
CREDENCIAIS = [
    ("planilhas (PAINEL_TRIBO_SHEETS_SA)", os.environ.get("SHEETS_SA_FILE")),
    ("planilhas (GOOGLE_SHEETS_SA_JSON)", os.environ.get("SHEETS_SA_ALT_FILE")),
    ("BigQuery (GCP_SA_KEY)", os.environ.get("GOOGLE_APPLICATION_CREDENTIALS")),
    ("arquivo local", str(ROOT / "google_sa.json")),
]
SHEET_ID = os.environ.get("MENS_SHEET_ID") or "1gfo0ORs4ccD0yn13eCuNxHM9USoGH_vjuzytcS1-QyQ"
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly",
          "https://www.googleapis.com/auth/drive.readonly"]


def tokens_disponiveis():
    """(apelido, token) de cada credencial que existe no ambiente."""
    from google.oauth2 import service_account
    import google.auth.transport.requests
    for apelido, caminho in CREDENCIAIS:
        if not caminho or not Path(caminho).exists():
            continue
        try:
            creds = service_account.Credentials.from_service_account_file(caminho, scopes=SCOPES)
            creds.refresh(google.auth.transport.requests.Request())
        except Exception as e:
            print(f"[mens] credencial {apelido} não carregou: {type(e).__name__}")
            continue
        # O e-mail da SA NÃO vai para o log: este repositório é público.
        yield apelido, creds.token


def baixar_xlsx() -> bytes:
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    cab = {"User-Agent": "Mozilla/5.0"}
    try:
        req = urllib.request.Request(url, headers=cab)
        with urllib.request.urlopen(req, timeout=120) as r:
            dados = r.read()
        if dados[:2] == b"PK":
            print("[mens] planilha baixada sem credencial (link público)")
            return dados
    except urllib.error.HTTPError as e:
        print(f"[mens] link público respondeu HTTP {e.code}; tentando com a service account")

    tentou = False
    for apelido, token in tokens_disponiveis():
        tentou = True
        cab["Authorization"] = f"Bearer {token}"
        try:
            req = urllib.request.Request(url, headers=cab)
            with urllib.request.urlopen(req, timeout=120) as r:
                dados = r.read()
        except urllib.error.HTTPError as e:
            print(f"[mens] credencial {apelido}: HTTP {e.code}")
            continue
        if dados[:2] == b"PK":
            print(f"[mens] planilha baixada com a credencial {apelido}")
            return dados
        print(f"[mens] credencial {apelido}: a resposta não é um xlsx")
    if not tentou:
        raise RuntimeError("nenhuma credencial de service account encontrada no ambiente")
    raise RuntimeError("nenhuma credencial tem acesso à planilha — compartilhe-a como Leitor")


def main() -> None:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(baixar_xlsx()), data_only=True, read_only=True)
    print(f"\nABAS ({len(wb.sheetnames)}):")
    for n in wb.sheetnames:
        print("  " + n)

    for nome in wb.sheetnames:
        ws = wb[nome]
        linhas = 0
        cabecalho: list[str] = []
        for i, l in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                cabecalho = [str(c).strip() if c is not None else "" for c in l]
            if l and any(c not in (None, "") for c in l):
                linhas += 1
            if i > 4000:
                break
        print(f"\n=== {nome} === {linhas} linhas com conteúdo")
        for i, c in enumerate(cabecalho):
            if c:
                print(f"   [{i}] {c}")
        # Uma amostra do TIPO de cada coluna, nunca do valor.
        amostra = None
        for i, l in enumerate(ws.iter_rows(min_row=2, max_row=6, values_only=True)):
            if l and any(c not in (None, "") for c in l):
                amostra = l
                break
        if amostra:
            tipos = [type(c).__name__ if c is not None else "-" for c in amostra[:24]]
            print("   tipos da 1ª linha: " + " ".join(tipos))


if __name__ == "__main__":
    main()
