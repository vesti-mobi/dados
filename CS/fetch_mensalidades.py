"""
fetch_mensalidades.py — o que cada marca paga, da planilha de mensalidades.

Pedido da Laura em 29/10/2026: "em bonificações aparecem marcas que não têm
valor porque não têm mensalidade na Iugu; quando acontecer isso, puxe daqui
[planilha], sempre da aba vesti e starter e sempre olhando no mês mais atual
que tem lá".

O caso que originou: o upgrade da Boise (R$ 1.299) entrava no quadrante de
Cross-sell com ganho ZERO. Upgrade conta pela DIFERENÇA de mensalidade, e a
Boise não tem mensalidade nenhuma no Iugu — 293 dias de série, tudo zero,
nenhuma fatura. Sem base anterior não há diferença, e o valor sumia. São
quatro marcas assim só no trimestre.

A planilha tem uma aba por mês e por produto — "Vesti 11-2026",
"Starter 11-2026" —, e os nomes não seguem um padrão só: há "Vesti 1225",
"Starter 0126", "Vesti 072026". `mes_da_aba` lê todos.

Saída: CS/mensalidades_planilha.json, intermediário, fora do git.

Acesso: a planilha NÃO é pública (o export responde 401). Vai pela service
account de leitura de planilhas, cuja chave o workflow grava a partir de um
secret. O e-mail dela nunca aparece no log: este repositório é público.

Rodar:  python3 fetch_mensalidades.py
"""

from __future__ import annotations

import io
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(__file__).parent
OUT = ROOT / "mensalidades_planilha.json"
SHEET_ID = os.environ.get("MENS_SHEET_ID") or "1gfo0ORs4ccD0yn13eCuNxHM9USoGH_vjuzytcS1-QyQ"
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly",
          "https://www.googleapis.com/auth/drive.readonly"]

# Em ordem de preferência. A planilha foi compartilhada com a service account
# de leitura de planilhas, que não é a do BigQuery.
CREDENCIAIS = [
    ("planilhas", os.environ.get("SHEETS_SA_FILE")),
    ("planilhas (alt)", os.environ.get("SHEETS_SA_ALT_FILE")),
    ("BigQuery", os.environ.get("GOOGLE_APPLICATION_CREDENTIALS")),
    ("arquivo local", str(ROOT / "google_sa.json")),
]

# "sempre da aba vesti e starter" (Laura). As outras abas da planilha —
# "entradas 01-02-0325" e afins — ficam de fora.
PRODUTOS = ("vesti", "starter")


def chave(s) -> str:
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def numero(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = re.sub(r"[^\d,.-]", "", str(v)).replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return 0.0


def mes_da_aba(nome: str) -> tuple[str, str] | None:
    """('vesti'|'starter', 'AAAA-MM') ou None.

    Os nomes não têm um formato só: "Vesti 11-2026", "Vesti 1225",
    "Starter 0126", "Vesti 072026", "Starter 09-25". O que todos têm é mês e
    ano grudados na ordem MMAAAA ou MMAA, com ou sem separador — então a
    leitura é pelos DÍGITOS, não pelo desenho do nome.
    """
    k = chave(nome)
    produto = next((p for p in PRODUTOS if k.startswith(p)), None)
    if not produto:
        return None
    digitos = re.sub(r"\D", "", k)
    if len(digitos) == 6:
        mm, aaaa = digitos[:2], digitos[2:]
    elif len(digitos) == 4:
        mm, aaaa = digitos[:2], "20" + digitos[2:]
    else:
        return None
    if not ("01" <= mm <= "12") or not ("2020" <= aaaa <= "2099"):
        return None
    return produto, f"{aaaa}-{mm}"


def baixar_xlsx() -> bytes:
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    cab = {"User-Agent": "Mozilla/5.0"}
    try:
        with urllib.request.urlopen(urllib.request.Request(url, headers=cab), timeout=120) as r:
            dados = r.read()
        if dados[:2] == b"PK":
            print("[mens] planilha baixada sem credencial (link público)")
            return dados
    except urllib.error.HTTPError as e:
        print(f"[mens] link público respondeu HTTP {e.code}; tentando com a service account")

    from google.oauth2 import service_account
    import google.auth.transport.requests

    tentou = False
    for apelido, caminho in CREDENCIAIS:
        if not caminho or not Path(caminho).exists():
            continue
        tentou = True
        try:
            creds = service_account.Credentials.from_service_account_file(caminho, scopes=SCOPES)
            creds.refresh(google.auth.transport.requests.Request())
            # O e-mail da SA NÃO vai para o log: este repositório é público.
            c = dict(cab, Authorization=f"Bearer {creds.token}")
            with urllib.request.urlopen(urllib.request.Request(url, headers=c), timeout=120) as r:
                dados = r.read()
        except Exception as e:
            print(f"[mens] credencial {apelido}: {type(e).__name__}")
            continue
        if dados[:2] == b"PK":
            print(f"[mens] planilha baixada com a credencial {apelido}")
            return dados
    if not tentou:
        raise RuntimeError("nenhuma credencial de service account encontrada")
    raise RuntimeError("nenhuma credencial tem acesso — compartilhe a planilha como Leitor")


def main() -> None:
    import openpyxl

    try:
        xlsx = baixar_xlsx()
    except Exception as e:
        print(f"[mens] {e}")
        print("[mens] sem planilha — o painel segue sem o valor de quem não tem mensalidade no Iugu")
        return

    wb = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True, read_only=True)

    # "sempre olhando no mês mais atual que tem lá" — por produto, porque as
    # duas famílias nem sempre fecham o mês no mesmo dia.
    melhor: dict[str, tuple[str, str]] = {}   # produto -> (mês, nome da aba)
    for nome in wb.sheetnames:
        r = mes_da_aba(nome)
        if not r:
            continue
        produto, mes = r
        if produto not in melhor or mes > melhor[produto][0]:
            melhor[produto] = (mes, nome)
    if not melhor:
        print("[mens] nenhuma aba de Vesti ou Starter reconhecida")
        return

    marcas: dict[str, dict] = {}
    for produto, (mes, aba) in sorted(melhor.items()):
        ws = wb[aba]
        linhas = list(ws.iter_rows(values_only=True))
        if not linhas:
            continue
        cab = [chave(c) for c in linhas[0]]
        idx = {c: i for i, c in enumerate(cab) if c}

        def pegar(l, *nomes):
            for n in nomes:
                i = idx.get(n)
                if i is not None and i < len(l) and l[i] not in (None, ""):
                    return l[i]
            return None

        n = 0
        for l in linhas[1:]:
            if not l or all(c in (None, "") for c in l):
                continue
            dom = re.sub(r"\D", "", str(pegar(l, "dominio", "domínio", "domain") or ""))
            nome_marca = str(pegar(l, "marca", "company", "cliente") or "").strip()
            if not dom and not nome_marca:
                continue
            # MENSALIDADE é só o plano; TOTAL COBRADO soma integração,
            # assistente e filial. Para medir o ganho de um UPGRADE DE PLANO a
            # base justa é plano contra plano — por isso `mensalidade` é o que
            # o painel usa, e `total` fica ao lado para quem precisar do que a
            # marca paga por inteiro (é o que o churn usa).
            reg = {
                "dominio": dom or None,
                "marca": nome_marca or None,
                "plano": (str(pegar(l, "plano") or "")).strip() or None,
                "mensalidade": numero(pegar(l, "mensalidade")),
                "total": numero(pegar(l, "total cobrado")),
                "cnpj": re.sub(r"\D", "", str(pegar(l, "cpfcnpj", "cpf e cnpj", "cnpj") or "")) or None,
                "produto": produto,
                "mes": mes,
            }
            if not reg["mensalidade"] and not reg["total"]:
                continue
            # UMA MARCA PODE TER VÁRIAS LINHAS. A Villon tem três na
            # planilha — a matriz e duas filiais —, e o que ela paga é a soma.
            # Então o registro guarda as duas leituras:
            #   `mensalidade` e `total` são da LINHA PRINCIPAL (a de maior
            #     mensalidade), que é o plano da marca — é o que serve de base
            #     para medir a diferença de um upgrade de plano;
            #   `totalGeral` e `mensalidadeGeral` somam TODAS as linhas, que é
            #     o que a marca paga por inteiro.
            # Misturar os dois daria base errada no upgrade: comparar o plano
            # novo de uma loja com a soma da matriz e das filiais.
            k = ("d:" + dom) if dom else ("n:" + chave(nome_marca))
            antes = marcas.get(k)
            if not antes:
                reg["linhas"] = 1
                reg["mensalidadeGeral"] = reg["mensalidade"]
                reg["totalGeral"] = reg["total"]
                marcas[k] = reg
            else:
                antes["linhas"] += 1
                antes["mensalidadeGeral"] += reg["mensalidade"]
                antes["totalGeral"] += reg["total"]
                if reg["mensalidade"] > antes["mensalidade"]:
                    antes["mensalidade"] = reg["mensalidade"]
                    antes["total"] = reg["total"]
                    antes["plano"] = reg["plano"]
                    antes["marca"] = reg["marca"] or antes["marca"]
            n += 1
        print(f'[mens] aba "{aba}" ({mes}): {n} linhas com valor')

    saida = {
        "geradoEm": datetime.now().astimezone().isoformat(timespec="seconds"),
        "abas": {p: m[1] for p, m in sorted(melhor.items())},
        "meses": {p: m[0] for p, m in sorted(melhor.items())},
        "marcas": marcas,
    }
    OUT.write_text(json.dumps(saida, ensure_ascii=False), encoding="utf-8")
    varias = sum(1 for r in marcas.values() if r.get("linhas", 1) > 1)
    print(f"[mens] {len(marcas)} marcas com mensalidade"
          + (f" · {varias} com mais de uma linha (matriz + filiais, somadas)" if varias else ""))
    print(f"[write] {OUT.name}")


if __name__ == "__main__":
    main()
